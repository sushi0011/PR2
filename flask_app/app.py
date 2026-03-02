from flask import Flask, render_template, request, jsonify, send_file
import os
import sqlite3
import json
from datetime import datetime, date
import uuid
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import io

app = Flask(__name__)

# ─── DATABASE ─────────────────────────────────────────────────────────────────
DB_PATH = os.path.join(os.path.dirname(__file__), "data", "pr_data.db")

def get_db():
    os.makedirs(os.path.dirname(DB_PATH), exist_ok=True)
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA foreign_keys = ON")
    return conn

def init_db():
    conn = get_db()
    c = conn.cursor()
    c.execute("""
        CREATE TABLE IF NOT EXISTS pr (
            id          TEXT PRIMARY KEY,
            number      TEXT NOT NULL,
            title       TEXT NOT NULL,
            category    TEXT NOT NULL,
            status      TEXT NOT NULL DEFAULT 'en-cours',
            created_date TEXT NOT NULL
        )
    """)
    c.execute("""
        CREATE TABLE IF NOT EXISTS task (
            pr_id       TEXT NOT NULL,
            task_id     TEXT NOT NULL,
            title       TEXT,
            description TEXT,
            done        INTEGER NOT NULL DEFAULT 0,
            date_prev   TEXT DEFAULT '',
            date_reelle TEXT DEFAULT '',
            note        TEXT DEFAULT '',
            PRIMARY KEY (pr_id, task_id),
            FOREIGN KEY (pr_id) REFERENCES pr(id) ON DELETE CASCADE
        )
    """)
    # Migration: add date_prev / date_reelle columns if they don't exist yet
    existing = [row[1] for row in c.execute("PRAGMA table_info(task)").fetchall()]
    if "date_prev" not in existing:
        c.execute("ALTER TABLE task ADD COLUMN date_prev TEXT DEFAULT ''")
    if "date_reelle" not in existing:
        c.execute("ALTER TABLE task ADD COLUMN date_reelle TEXT DEFAULT ''")
    
    c.execute("""
        CREATE TABLE IF NOT EXISTS document (
            id          TEXT PRIMARY KEY,
            name        TEXT NOT NULL,
            status      TEXT NOT NULL DEFAULT 'En attente',
            color       TEXT NOT NULL DEFAULT '#3498db',
            done        INTEGER NOT NULL DEFAULT 0,
            created_date TEXT NOT NULL
        )
    """)
    conn.commit()
    conn.close()

# ─── PROCESSING TIME LIMITS (in weeks) ────────────────────────────────────────
PROCESSING_LIMITS = {
    "ED": {
        "max_weeks": 10,
        "steps": {
            "Délai de dépot des offres": 3,
            "Etude technique": 2,
            "Négociation": 2,
            "Contractualisation": 3,
        }
    },
    "CR": {
        "max_weeks": 11,
        "steps": {
            "Délai de dépot des offres": 4,
            "Etude technique": 2,
            "Négociation": 2,
            "Contractualisation": 3,
        }
    },
    "COU": {
        "max_weeks": 14,
        "steps": {
            "Délai de dépot des offres": 5,
            "Etude technique": 3,
            "Négociation": 3,
            "Contractualisation": 3,
        }
    },
}

# ─── STEPS DEFINITIONS ────────────────────────────────────────────────────────
STEPS_DATA = {
    "ED": [
        {"id": 1,  "title": "Demande et vérification de la fiche de justification de l'ED",
         "desc": "Demander la fiche de justification de l'ED, dûment signée et cachetée par la direction du SU, et vérifier au préalable si elle est déjà disponible en pièces jointes sur Ariba."},
        {"id": 2,  "title": "Programmation du dossier en Commission de Marchés (CM)",
         "desc": "Programmer le dossier pour concertation en CM."},
        {"id": 3,  "title": "Suite à la concertation",
         "desc": "Après la concertation, obtenir l'OK pour faire aboutir le dossier ou poursuivre la procédure."},
        {"id": 4,  "title": "Rédaction et envoi du PV de la CM",
         "desc": "Rédiger le procès-verbal (PV) de la CM (pour concertation) et l'envoyer, pour signature, en réponse au mail de programmation de la CM, en plaçant le Chef de Division de la section CC en tête des destinataires."},
        {"id": 5,  "title": "Demande du BOQ et des spécifications techniques",
         "desc": "Envoyer au SU un mail, en réponse au mail de notification Ariba du PR, afin de demander le BOQ et les spécifications techniques."},
        {"id": 6,  "title": "Rédaction du PV de la première réunion de binôme",
         "desc": "Après réception de la réponse du SU, rédiger le procès-verbal de la première réunion de binôme."},
        {"id": 7,  "title": "Transmission du PV de la première réunion de binôme",
         "desc": "Envoyer le PV de la première réunion de binôme au SU, en mettant en copie la hiérarchie achats et la hiérarchie SU. (Hors directeurs)"},
        {"id": 8,  "title": "Réception des spécifications techniques",
         "desc": "Ajouter le BOQ comme tableau, enregistrer sous forme de PDF, joindre les conditions générales d'achat PDF et envoyer le mail à l'adresse du Fournisseur. N.B : Modifier le titre et les détails s'ils décrivent à quoi servent ces articles — le fournisseur n'a pas intérêt à savoir ça, juste les articles."},
        {"id": 9,  "title": "Réception des offres + PV d'ouverture technique",
         "desc": "Le fournisseur doit envoyer l'offre technique et commerciale avant ou le jour du délai. Transférer l'offre technique au service utilisateur pour étude de conformité technique. Le SU doit envoyer le rapport technique signé et cacheté + PV d'ouverture technique."},
        {"id": 10, "title": "Réception du PV d'OT et Rapport de conformité technique",
         "desc": "Après réception, envoyer le mail de l'offre commerciale du fournisseur au SU. Ensuite, rédiger le PV d'ouverture de l'OI et l'envoyer également au SU."},
        {"id": 11, "title": "Évaluation OI et relance pour OA1",
         "desc": "Réaliser une évaluation commerciale de l'OI (vérifier que les calculs du fournisseur sont cohérents avec les nôtres) et relancer le fournisseur pour OA1."},
        {"id": 12, "title": "Évaluation OA1 et relance pour OA2",
         "desc": "Transférer le mail de l'OA1 au SU, et relancer le fournisseur pour OA2."},
        {"id": 13, "title": "Évaluation OA2 et relance pour OA3",
         "desc": "Transférer le mail de l'OA2 au SU, et relancer le fournisseur pour OA3 en lui demandant de s'aligner sur le prix historique/Budget (le montant le moins cher entre ces deux)."},
        {"id": 14, "title": "Réception de l'offre finale et fin des négociations commerciales",
         "desc": "Demander les raisons de non-alignement si c'est le cas, conclure les négociations commerciales et envoyer l'OA3 au SU."},
        {"id": 15, "title": "Rédiger le Rapport de Binôme",
         "desc": "Rédiger le rapport de binôme, envoyer pour signature au SU et par notre direction avant CM."},
        {"id": 16, "title": "Planifier en CM et Adjudication",
         "desc": "Planifier le dossier par adjudication avec Rapport de binôme en PJ par mail, envoyer au chef de division + Planifier sur le CMCO dans NAS."},
        {"id": 17, "title": "Adjudication et rédaction du rapport de CM",
         "desc": "Rédiger le rapport de CM pour adjudication de ce dossier."},
        {"id": 18, "title": "Fiche de décision d'ED",
         "desc": "Fiche de décision d'entente directe, à ajouter dans le dossier."},
        {"id": 19, "title": "Attente de l'OK et Création du BC sur Ariba",
         "desc": "Après réception de l'OK Vert par mail, l'imprimer et garder dans le dossier, et créer le BC sur Ariba pour être envoyé au fournisseur et clôturer le dossier."},
    ],
    "CR": [
        {"id": 1,  "title": "Réception de la PR sur Ariba",
         "desc": "Vérifier que la PR soit une CR et contrôler les pièces jointes (comme le CPS)."},
        {"id": 2,  "title": "Demander le BOQ et les spécifications techniques",
         "desc": "Si ces deux documents ne sont pas joints sur Ariba, les demander par mail au SU et à sa hiérarchie."},
        {"id": 3,  "title": "Rédiger le PV1 de binôme",
         "desc": "Rédiger, vérifier et envoyer le PV1 au SU pour signature et retour."},
        {"id": 4,  "title": "Lire, vérifier et modifier le CPS — le transmettre pour avis du SU",
         "desc": "Bien lire le CPS et vérifier le contenu, ajouter la page de garde et les commentaires nécessaires au SU, puis le transmettre pour avis et retour."},
        {"id": 5,  "title": "Envoi du CPS final pour validation juridique + Réception PV1",
         "desc": "Dépôt du CPS sur la plateforme 'Galaxy' pour validation juridique et réception du PV1 signé."},
        {"id": 6,  "title": "Lancer la consultation aux fournisseurs du panel",
         "desc": "Par emails, en utilisant la copie cachée (Bcc), avec la hiérarchie en copie."},
        {"id": 7,  "title": "Dépôt des dossiers par mails séparés + Demande des raisons de non-participation",
         "desc": "Demander aux fournisseurs non-soumissionnaires leurs raisons de non-participation. Justifier si on a des adresses mail erronées."},
        {"id": 8,  "title": "Transmettre les raisons de non participation au SU et avoir leur OK pour continuer le processus",
         "desc": "Mail standard et joigné les réponses de non participations (copier les mails des sociétés)."},
        {"id": 9,  "title": "Envoyer les Mails techniques des sociétés soumissionnaires au SU et demander d'accusé leur récéption",
         "desc": "Demander l'accusé de récéption du SU puis Demander les mots de passes techniques aux soumissionnaires."},
        {"id": 10, "title": "Transferer les mails des offres techniques au SU (Avec Mots de passes)",
         "desc": "Envoyer les mails des offres techniques en PJ avec les Mots de passes communiqués au SU."},
        {"id": 11, "title": "Rédiger le PV d'OT + Demander le Rapport de conformité technique",
         "desc": "PV d'OT a envoyer et attente du rapport de conformité technique du SU."},
        {"id": 12, "title": "Ecarter les societés Non Conformes selon le Rapport + Envoyer les offeres commerciales",
         "desc": "Envoyer les offres financiers au SU des societés conformes et rédiger le PV d'Ouverture et Evalution commerciale."},
        {"id": 13, "title": "Evaluer les OI + Demander OA1 et evaluer selon la procédure",
         "desc": "Relancer les 3 premier moins disant + le 4ème si il presente un écart < 15% ."},
        {"id": 14, "title": "Evaluer les OA1 + Demander OA2 et evaluer selon la procédure",
         "desc": "Relancer le premier moins disant + le 2ème si il presente un écart < 5% ."},
        {"id": 15, "title": "Evaluer les OA2 + Demander OA3",
         "desc": "Conclure les négociations commerciales et relancer seulement le moins disant."},
        {"id": 16, "title": "Rédiger le Rapport de Binôme",
         "desc": "Rédiger le rapport de binôme final, envoyer pour signature au SU et à notre direction avant CM."},
        {"id": 17, "title": "Planifier en CM et Adjudication",
         "desc": "Planifier le dossier par adjudication avec rapport de binôme en PJ, envoyer au chef de division et planifier sur le CMCO dans NAS."},
        {"id": 18, "title": "Adjudication et rédaction du rapport de CM",
         "desc": "Rédiger le rapport de CM pour adjudication du dossier."},
        {"id": 19, "title": "Attente de l'OK et Création du BC sur Ariba",
         "desc": "Après réception de l'OK Vert par mail, l'imprimer et garder dans le dossier, et créer le BC sur Ariba pour être envoyé au fournisseur et clôturer le dossier."},
        {"id": 20, "title": "Clôture et archivage du dossier",
         "desc": "Archiver l'ensemble des documents du dossier (PV, rapports, offres, BC) et clôturer le dossier sur Ariba."},
    ],
    "COU": [
        {"id": 1, "title": "Publication de l'avis d'appel d'offres",
         "desc": "Publier l'avis d'appel d'offres sur les plateformes officielles et dans les journaux habilités conformément à la réglementation en vigueur."},
        {"id": 2, "title": "Période de consultation",
         "desc": "Gérer la période de consultation des offres : répondre aux questions des candidats, diffuser les éventuels amendements au CPS."},
        {"id": 3, "title": "Réception et traitement des offres",
         "desc": "Collecter les offres sous pli fermé, les enregistrer et préparer la séance d'ouverture avec le PV correspondant."},
        {"id": 4, "title": "Attribution et notification",
         "desc": "Attribuer le marché au soumissionnaire retenu, notifier officiellement tous les soumissionnaires du résultat et créer le BC sur Ariba pour clôturer le dossier."},
    ],
}

# ─── HELPERS ──────────────────────────────────────────────────────────────────
def compute_progress(tasks: dict) -> int:
    if not tasks:
        return 0
    completed = sum(1 for t in tasks.values() if t.get("done"))
    return round((completed / len(tasks)) * 100)

def delay_status(date_prev: str, date_reelle: str) -> str:
    """
    Returns:
      'ontime'  — réelle <= prév
      'warning' — réelle entre 1 et 5 jours après prév
      'late'    — réelle > 5 jours après prév
      ''        — données manquantes
    """
    if not date_prev or not date_reelle:
        return ""
    try:
        dp = date.fromisoformat(date_prev)
        dr = date.fromisoformat(date_reelle)
        delta = (dr - dp).days
        if delta <= 0:
            return "ontime"
        elif delta <= 5:
            return "warning"
        else:
            return "late"
    except ValueError:
        return ""

def calculate_processing_time(tasks: dict, created_date: str) -> dict:
    """
    Calculate total processing time in days from created_date to last completed step.
    Returns: { "days": int, "weeks": float, "exceeded": bool, "max_weeks": int }
    """
    if not tasks or not created_date:
        return {"days": 0, "weeks": 0, "exceeded": False, "max_weeks": 0}
    
    # Find last task with date_reelle
    last_completion = None
    for task in tasks.values():
        if task.get("date_reelle"):
            try:
                dr = date.fromisoformat(task["date_reelle"])
                if not last_completion or dr > last_completion:
                    last_completion = dr
            except ValueError:
                continue
    
    if not last_completion:
        return {"days": 0, "weeks": 0, "exceeded": False, "max_weeks": 0}
    
    try:
        start_date = date.fromisoformat(created_date)
        total_days = (last_completion - start_date).days
        total_weeks = round(total_days / 7, 1)
        return {
            "days": total_days,
            "weeks": total_weeks,
            "exceeded": False,  # Will be set by API based on category
            "max_weeks": 0
        }
    except ValueError:
        return {"days": 0, "weeks": 0, "exceeded": False, "max_weeks": 0}

def row_to_pr(row) -> dict:
    return {
        "id":          row["id"],
        "number":      row["number"],
        "title":       row["title"],
        "category":    row["category"],
        "status":      row["status"],
        "createdDate": row["created_date"],
    }

def load_tasks(conn, pr_id: str) -> dict:
    rows = conn.execute(
        "SELECT * FROM task WHERE pr_id = ? ORDER BY CAST(task_id AS INTEGER)",
        (pr_id,)
    ).fetchall()
    return {
        str(r["task_id"]): {
            "title":       r["title"] or "",
            "desc":        r["description"] or "",
            "done":        bool(r["done"]),
            "date_prev":   r["date_prev"] or "",
            "date_reelle": r["date_reelle"] or "",
            "note":        r["note"] or "",
            "delay":       delay_status(r["date_prev"], r["date_reelle"]),
        }
        for r in rows
    }

def count_late_steps_for_pr(tasks: dict) -> dict:
    """Returns counts of warning and late steps."""
    warning = sum(1 for t in tasks.values() if t.get("delay") == "warning")
    late    = sum(1 for t in tasks.values() if t.get("delay") == "late")
    return {"warning": warning, "late": late}

# ─── ROUTES ───────────────────────────────────────────────────────────────────

@app.route("/")
def index():
    conn = get_db()
    rows = conn.execute("SELECT * FROM pr").fetchall()
    total    = len(rows)
    cloturee = sum(1 for r in rows if r["status"] == "cloturee")
    en_cours = sum(1 for r in rows if r["status"] == "en-cours")
    blockee  = sum(1 for r in rows if r["status"] == "blockee")
    annulee  = sum(1 for r in rows if r["status"] == "annulee")

    total_prog = 0
    total_late = 0
    total_warning = 0
    for r in rows:
        tasks = load_tasks(conn, r["id"])
        total_prog += compute_progress(tasks)
        counts = count_late_steps_for_pr(tasks)
        total_late    += counts["late"]
        total_warning += counts["warning"]

    avg_prog = round(total_prog / total, 1) if total else 0
    conn.close()

    stats = dict(
        total=total, cloturee=cloturee, en_cours=en_cours,
        blockee=blockee, annulee=annulee, avg_prog=avg_prog,
        total_late=total_late, total_warning=total_warning
    )
    return render_template("index.html", stats=stats)


# ── PR CRUD ──────────────��────────────────────────────────────────────────────

@app.route("/api/pr", methods=["GET"])
def get_all_pr():
    conn = get_db()
    q = request.args.get("q", "").lower()
    rows = conn.execute("SELECT * FROM pr ORDER BY created_date DESC").fetchall()
    result = []
    for r in rows:
        if q and q not in str(r["number"]).lower() \
             and q not in r["title"].lower() \
             and q not in r["category"].lower():
            continue
        pr = row_to_pr(r)
        tasks = load_tasks(conn, r["id"])
        pr["progress"]        = compute_progress(tasks)
        pr["completed_tasks"] = sum(1 for t in tasks.values() if t.get("done"))
        pr["total_tasks"]     = len(tasks)
        counts = count_late_steps_for_pr(tasks)
        pr["late_steps"]    = counts["late"]
        pr["warning_steps"] = counts["warning"]
        
        # Add processing time info
        proc_time = calculate_processing_time(tasks, r["created_date"])
        pr["processing_days"]  = proc_time["days"]
        pr["processing_weeks"] = proc_time["weeks"]
        max_weeks = PROCESSING_LIMITS.get(r["category"], {}).get("max_weeks", 0)
        pr["max_weeks"] = max_weeks
        pr["exceeded"] = proc_time["weeks"] > max_weeks if max_weeks > 0 else False
        
        result.append(pr)
    conn.close()
    return jsonify(result)


@app.route("/api/pr", methods=["POST"])
def create_pr():
    body     = request.json
    number   = body.get("number","").strip()
    title    = body.get("title","").strip()
    category = body.get("category","")
    pr_date  = body.get("prDate","")

    if not all([number, title, category, pr_date]):
        return jsonify({"error": "Tous les champs sont requis"}), 400
    if category not in STEPS_DATA:
        return jsonify({"error": "Catégorie invalide"}), 400

    pr_id = f"PR-{number}-{datetime.now().strftime('%f')}"
    conn  = get_db()
    conn.execute(
        "INSERT INTO pr (id, number, title, category, status, created_date) VALUES (?,?,?,?,?,?)",
        (pr_id, number, title, category, "en-cours", pr_date)
    )
    for s in STEPS_DATA[category]:
        conn.execute(
            "INSERT INTO task (pr_id, task_id, title, description, done, date_prev, date_reelle, note) VALUES (?,?,?,?,0,'','','')",
            (pr_id, str(s["id"]), s["title"], s["desc"])
        )
    conn.commit()
    conn.close()
    return jsonify({"id": pr_id, "message": "PR créée avec succès"}), 201


@app.route("/api/pr/<pr_id>", methods=["GET"])
def get_pr(pr_id):
    conn = get_db()
    row = conn.execute("SELECT * FROM pr WHERE id = ?", (pr_id,)).fetchone()
    if not row:
        conn.close()
        return jsonify({"error": "PR introuvable"}), 404
    pr = row_to_pr(row)
    tasks = load_tasks(conn, pr_id)
    pr["tasks"]    = tasks
    pr["progress"] = compute_progress(tasks)
    conn.close()
    return jsonify(pr)


@app.route("/api/pr/<pr_id>", methods=["PUT"])
def update_pr(pr_id):
    conn = get_db()
    row = conn.execute("SELECT id FROM pr WHERE id = ?", (pr_id,)).fetchone()
    if not row:
        conn.close()
        return jsonify({"error": "PR introuvable"}), 404
    body = request.json
    fields = []
    values = []
    mapping = {"number":"number","title":"title","category":"category",
               "status":"status","createdDate":"created_date"}
    for k, col in mapping.items():
        if k in body:
            fields.append(f"{col} = ?")
            values.append(body[k])
    if fields:
        values.append(pr_id)
        conn.execute(f"UPDATE pr SET {', '.join(fields)} WHERE id = ?", values)
        conn.commit()
    conn.close()
    return jsonify({"message": "PR mise à jour"})


@app.route("/api/pr/<pr_id>", methods=["DELETE"])
def delete_pr(pr_id):
    conn = get_db()
    row = conn.execute("SELECT id FROM pr WHERE id = ?", (pr_id,)).fetchone()
    if not row:
        conn.close()
        return jsonify({"error": "PR introuvable"}), 404
    conn.execute("DELETE FROM task WHERE pr_id = ?", (pr_id,))
    conn.execute("DELETE FROM pr WHERE id = ?",      (pr_id,))
    conn.commit()
    conn.close()
    return jsonify({"message": "PR supprimée"})


# ── TASK (checklist) ──────────────────────────────────────────────────────────

@app.route("/api/pr/<pr_id>/task/<task_id>", methods=["PATCH"])
def update_task(pr_id, task_id):
    conn = get_db()
    row = conn.execute(
        "SELECT * FROM task WHERE pr_id = ? AND task_id = ?", (pr_id, task_id)
    ).fetchone()
    if not row:
        conn.close()
        return jsonify({"error": "Tâche introuvable"}), 404

    body   = request.json
    fields = []
    values = []
    allowed = {"done": "done", "date_prev": "date_prev",
               "date_reelle": "date_reelle", "note": "note"}
    for k, col in allowed.items():
        if k in body:
            val = body[k]
            if col == "done":
                val = 1 if val else 0
            fields.append(f"{col} = ?")
            values.append(val)
    if fields:
        values += [pr_id, task_id]
        conn.execute(f"UPDATE task SET {', '.join(fields)} WHERE pr_id = ? AND task_id = ?", values)
        conn.commit()

    tasks    = load_tasks(conn, pr_id)
    progress = compute_progress(tasks)
    counts   = count_late_steps_for_pr(tasks)
    # Return the updated task delay status too
    updated_task = tasks.get(task_id, {})
    conn.close()
    return jsonify({
        "progress":      progress,
        "late_steps":    counts["late"],
        "warning_steps": counts["warning"],
        "delay":         updated_task.get("delay", ""),
    })


# ── STATUS ────────────────────────────────────────────────────────────────────

@app.route("/api/pr/<pr_id>/status", methods=["PATCH"])
def set_status(pr_id):
    conn  = get_db()
    row   = conn.execute("SELECT id FROM pr WHERE id = ?", (pr_id,)).fetchone()
    if not row:
        conn.close()
        return jsonify({"error": "PR introuvable"}), 404
    status = request.json.get("status")
    valid  = ["en-cours", "cloturee", "blockee", "annulee"]
    if status not in valid:
        conn.close()
        return jsonify({"error": "Statut invalide"}), 400
    conn.execute("UPDATE pr SET status = ? WHERE id = ?", (status, pr_id))
    conn.commit()
    conn.close()
    return jsonify({"message": "Statut mis à jour"})


# ── STEPS REFERENCE ───────────────────────────────────────────────────────────

@app.route("/api/steps/<category>")
def get_steps(category):
    if category not in STEPS_DATA:
        return jsonify({"error": "Catégorie invalide"}), 404
    return jsonify(STEPS_DATA[category])


# ── ALERTS API ────────────────────────────────────────────────────────────────

@app.route("/api/alerts")
def get_alerts():
    """Returns all PR/steps that are late or in warning state."""
    conn  = get_db()
    rows  = conn.execute("SELECT * FROM pr WHERE status = 'en-cours'").fetchall()
    alerts = []
    for r in rows:
        tasks  = load_tasks(conn, r["id"])
        for tid, task in tasks.items():
            if task["delay"] in ("warning", "late") and not task["done"]:
                alerts.append({
                    "pr_id":       r["id"],
                    "pr_number":   r["number"],
                    "pr_title":    r["title"],
                    "task_id":     tid,
                    "task_title":  task["title"],
                    "date_prev":   task["date_prev"],
                    "date_reelle": task["date_reelle"],
                    "delay":       task["delay"],
                })
    conn.close()
    alerts.sort(key=lambda a: (0 if a["delay"] == "late" else 1))
    return jsonify(alerts)


# ── EXPORT EXCEL ──────────────────────────────────────────────────────────────

@app.route("/api/export")
def export_excel():
    conn      = get_db()
    date_from = request.args.get("from", "")
    date_to   = request.args.get("to",   "")

    wb        = openpyxl.Workbook()
    ws_ov     = wb.active
    ws_ov.title = "Vue d'ensemble"

    red_fill    = PatternFill("solid", fgColor="C0392B")
    grey_fill   = PatternFill("solid", fgColor="F2F2F2")
    white_fill  = PatternFill("solid", fgColor="FFFFFF")
    green_fill  = PatternFill("solid", fgColor="D5F5E3")
    orange_fill = PatternFill("solid", fgColor="FDEBD0")
    late_fill   = PatternFill("solid", fgColor="FADBD8")
    header_font = Font(bold=True, color="FFFFFF", size=11)
    normal_font = Font(size=10)
    center      = Alignment(horizontal="center", vertical="center")
    thin        = Side(style="thin", color="DDDDDD")
    border      = Border(left=thin, right=thin, top=thin, bottom=thin)

    headers_ov = ["N° PR", "Titre", "Catégorie", "Statut", "Date demande",
                  "Progression (%)", "Étapes complétées", "Total étapes",
                  "Étapes en retard", "Étapes à risque"]
    for col, h in enumerate(headers_ov, 1):
        cell = ws_ov.cell(row=1, column=col, value=h)
        cell.font = header_font; cell.fill = red_fill
        cell.alignment = center; cell.border = border

    row_idx = 2
    rows = conn.execute("SELECT * FROM pr ORDER BY created_date DESC").fetchall()
    for r in rows:
        if date_from and r["created_date"] < date_from: continue
        if date_to   and r["created_date"] > date_to:   continue
        tasks     = load_tasks(conn, r["id"])
        completed = sum(1 for t in tasks.values() if t.get("done"))
        progress  = compute_progress(tasks)
        counts    = count_late_steps_for_pr(tasks)
        fill      = grey_fill if row_idx % 2 == 0 else white_fill
        row_data  = [r["number"], r["title"], r["category"], r["status"],
                     r["created_date"], progress, completed, len(tasks),
                     counts["late"], counts["warning"]]
        for col, val in enumerate(row_data, 1):
            cell = ws_ov.cell(row=row_idx, column=col, value=val)
            cell.font = normal_font; cell.fill = fill
            cell.alignment = center if col != 2 else Alignment(vertical="center")
            cell.border = border
        row_idx += 1

    col_widths_ov = [12, 35, 14, 14, 16, 18, 20, 14, 18, 18]
    for i, w in enumerate(col_widths_ov, 1):
        ws_ov.column_dimensions[get_column_letter(i)].width = w
    ws_ov.row_dimensions[1].height = 22

    # One sheet per PR with full task details
    for r in rows:
        if date_from and r["created_date"] < date_from: continue
        if date_to   and r["created_date"] > date_to:   continue
        sheet_name = f"PR-{r['number']}"[:31]
        ws = wb.create_sheet(title=sheet_name)
        headers_pr = ["N°", "Titre de l'étape", "Description",
                      "Complétée", "Date prévisionnelle", "Date réelle",
                      "Délai", "Notes"]
        for col, h in enumerate(headers_pr, 1):
            cell = ws.cell(row=1, column=col, value=h)
            cell.font = header_font; cell.fill = red_fill
            cell.alignment = center; cell.border = border

        tasks = load_tasks(conn, r["id"])
        for row_i, (tid, task) in enumerate(tasks.items(), 2):
            delay = task.get("delay", "")
            if delay == "ontime":
                row_fill = green_fill
            elif delay == "warning":
                row_fill = orange_fill
            elif delay == "late":
                row_fill = late_fill
            else:
                row_fill = grey_fill if row_i % 2 == 0 else white_fill

            delay_labels = {"ontime": "Dans les délais", "warning": "Risque retard",
                            "late": "En retard", "": "—"}
            row_data = [tid, task.get("title"), task.get("desc"),
                        "Oui" if task.get("done") else "Non",
                        task.get("date_prev",""), task.get("date_reelle",""),
                        delay_labels.get(delay, "—"), task.get("note","")]
            for col, val in enumerate(row_data, 1):
                cell = ws.cell(row=row_i, column=col, value=val)
                cell.font = normal_font; cell.fill = row_fill
                cell.border = border
                cell.alignment = center if col in [1, 4, 5, 6, 7] \
                    else Alignment(vertical="center", wrap_text=True)

        for i, w in enumerate([6, 32, 45, 12, 20, 20, 18, 40], 1):
            ws.column_dimensions[get_column_letter(i)].width = w

    conn.close()
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    filename = f"PR_Export_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    return send_file(output,
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                     as_attachment=True, download_name=filename)


# ── PROCESSING TIME LIMITS (API) ──────────────────────────────────────────────
@app.route("/api/processing-limits", methods=["GET"])
def get_processing_limits():
    return jsonify(PROCESSING_LIMITS)


# ── DOCUMENTS MANAGEMENT ──────────────────────────────────────────────────────

@app.route("/api/documents", methods=["GET"])
def get_documents():
    conn = get_db()
    c = conn.cursor()
    docs = c.execute("SELECT * FROM document ORDER BY created_date DESC").fetchall()
    conn.close()
    return jsonify([
        {
            "id": doc[0],
            "name": doc[1],
            "status": doc[2],
            "color": doc[3],
            "done": doc[4],
            "created_date": doc[5],
        }
        for doc in docs
    ])


@app.route("/api/documents", methods=["POST"])
def create_document():
    data = request.get_json()
    doc_id = str(uuid.uuid4())
    name = data.get("name", "").strip()
    status = data.get("status", "En attente").strip()
    color = data.get("color", "#3498db").strip()
    
    if not name:
        return jsonify({"error": "Document name is required"}), 400
    
    conn = get_db()
    c = conn.cursor()
    c.execute(
        "INSERT INTO document (id, name, status, color, done, created_date) VALUES (?, ?, ?, ?, 0, ?)",
        (doc_id, name, status, color, datetime.now().isoformat())
    )
    conn.commit()
    conn.close()
    
    return jsonify({
        "id": doc_id,
        "name": name,
        "status": status,
        "color": color,
        "done": 0,
        "created_date": datetime.now().isoformat(),
    }), 201


@app.route("/api/documents/<doc_id>", methods=["PUT"])
def update_document(doc_id):
    data = request.get_json()
    conn = get_db()
    c = conn.cursor()
    
    doc = c.execute("SELECT * FROM document WHERE id = ?", (doc_id,)).fetchone()
    if not doc:
        conn.close()
        return jsonify({"error": "Document not found"}), 404
    
    name = data.get("name", doc[1]).strip()
    status = data.get("status", doc[2]).strip()
    color = data.get("color", doc[3]).strip()
    done = data.get("done", doc[4])
    
    c.execute(
        "UPDATE document SET name = ?, status = ?, color = ?, done = ? WHERE id = ?",
        (name, status, color, done, doc_id)
    )
    conn.commit()
    conn.close()
    
    return jsonify({
        "id": doc_id,
        "name": name,
        "status": status,
        "color": color,
        "done": done,
        "created_date": doc[5],
    })


@app.route("/api/documents/<doc_id>", methods=["DELETE"])
def delete_document(doc_id):
    conn = get_db()
    c = conn.cursor()
    
    doc = c.execute("SELECT * FROM document WHERE id = ?", (doc_id,)).fetchone()
    if not doc:
        conn.close()
        return jsonify({"error": "Document not found"}), 404
    
    c.execute("DELETE FROM document WHERE id = ?", (doc_id,))
    conn.commit()
    conn.close()
    
    return jsonify({"success": True})


# ── IMPORT EXCEL ──────────────────────────────────────────────────────────────

@app.route("/api/import", methods=["POST"])
def import_excel():
    if "file" not in request.files:
        return jsonify({"error": "Aucun fichier fourni"}), 400
    f = request.files["file"]
    if not f.filename.endswith(".xlsx"):
        return jsonify({"error": "Format invalide — fichier .xlsx requis"}), 400
    try:
        wb = openpyxl.load_workbook(io.BytesIO(f.read()))
        if "Vue d'ensemble" not in wb.sheetnames and "Overview" not in wb.sheetnames:
            return jsonify({"error": "Onglet 'Vue d'ensemble' introuvable"}), 400
        ws_ov = wb["Vue d'ensemble"] if "Vue d'ensemble" in wb.sheetnames else wb["Overview"]
        conn = get_db()
        imported = 0
        for row in ws_ov.iter_rows(min_row=2, values_only=True):
            if not row[0]: continue
            number, title, category, status, created_date = \
                row[0], row[1], row[2], row[3], row[4]
            if category not in STEPS_DATA: continue
            pr_id      = f"PR-{number}-imported"
            sheet_name = f"PR-{number}"[:31]
            # upsert PR
            conn.execute(
                "INSERT OR REPLACE INTO pr (id, number, title, category, status, created_date) VALUES (?,?,?,?,?,?)",
                (pr_id, str(number), str(title), category,
                 status or "en-cours", str(created_date) if created_date else "")
            )
            conn.execute("DELETE FROM task WHERE pr_id = ?", (pr_id,))
            if sheet_name in wb.sheetnames:
                ws_pr = wb[sheet_name]
                for tr in ws_pr.iter_rows(min_row=2, values_only=True):
                    if not tr[0]: continue
                    tid, ttitle, tdesc, tdone, tprev, treelle, _delay, tnote = \
                        tr[0], tr[1], tr[2], tr[3], tr[4], tr[5], tr[6] if len(tr)>6 else "", tr[7] if len(tr)>7 else ""
                    conn.execute(
                        "INSERT INTO task (pr_id,task_id,title,description,done,date_prev,date_reelle,note) VALUES (?,?,?,?,?,?,?,?)",
                        (pr_id, str(tid), ttitle or "", tdesc or "",
                         1 if tdone == "Oui" else 0,
                         str(tprev) if tprev else "", str(treelle) if treelle else "",
                         tnote or "")
                    )
            else:
                for s in STEPS_DATA[category]:
                    conn.execute(
                        "INSERT INTO task (pr_id,task_id,title,description,done,date_prev,date_reelle,note) VALUES (?,?,?,?,0,'','','')",
                        (pr_id, str(s["id"]), s["title"], s["desc"])
                    )
            imported += 1
        conn.commit()
        conn.close()
        return jsonify({"message": f"{imported} PR importée(s) avec succès", "imported": imported})
    except Exception as e:
        return jsonify({"error": str(e)}), 500


# ─── MAIN ──────────────────────────────────────────────────────────────────────
if __name__ == "__main__":
    init_db()
    app.run(debug=True, port=5000)
